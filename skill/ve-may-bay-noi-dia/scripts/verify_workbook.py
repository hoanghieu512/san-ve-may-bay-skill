#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Đối chiếu độc lập workbook đã recalc với checkpoint. Bước 2 của SKILL.md §8.

    python3 verify_workbook.py checkpoint.json output.xlsx

Chạy SAU recalc.py. Tự tính lại mọi con số từ checkpoint bằng code riêng (không
dùng shortlist.py của build) rồi so với giá trị trong file:

- mọi ô công thức đã có giá trị và không lỗi (chưa recalc là bắt được ở đây);
- giá từng nguồn, Giá chốt, giá thấp nhất, nguồn giá thấp nhất, % chênh;
- ô `Không khả dụng` nào cũng có dòng Log giải thích;
- Block C: chiều đi/về rẻ nhất, cùng hãng rẻ nhất, tổng = đi + về;
- Summary: tổng khớp Block C, sắp tăng dần, đủ link hai chiều, tổng mua được rẻ nhất;
- Mua vé: giá khớp nguồn, đủ link + thời điểm, đủ chuyến đề xuất.

Thoát mã 1 nếu có lỗi. Recalc sạch chỉ chứng minh công thức chạy được, không chứng
minh đúng: script này mới là bước bắt lỗi.
"""
import datetime, json, sys
from openpyxl import load_workbook

NA = "Không khả dụng"
METASEARCH = {'Google Flights'}
LOI_EXCEL = ('#VALUE!', '#REF!', '#NAME?', '#DIV/0!', '#N/A', '#NUM!', '#NULL!')


class Kiem:
    def __init__(self):
        self.n, self.loi = 0, []

    def __call__(self, ok, msg):
        self.n += 1
        if not ok:
            self.loi.append(msg)
        return ok


def chot(f, src):
    return next((f['gia'][s] for s in src if f.get('gia', {}).get(s) is not None), None)


def bang(ws, hang_tieu_de):
    """{tên cột: số cột} của một dòng tiêu đề."""
    return {ws.cell(hang_tieu_de, c).value: c for c in range(1, ws.max_column + 1)
            if ws.cell(hang_tieu_de, c).value}


def dong_tieu_de(ws, nhan='Cặp ngày'):
    return [r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value == nhan]


def hhmm(v):
    return v.strftime('%H:%M') if isinstance(v, (datetime.time, datetime.datetime)) else v


def gan(a, b):
    return isinstance(a, (int, float)) and isinstance(b, (int, float)) and abs(a - b) < 0.5


def ban_cua(f, chang, src):
    """Giá ở người bán, tính độc lập: trang đặt vé GF (nếu OK) + nguồn quét là người bán."""
    out = []
    m = f.get('mua') or {}
    if m.get('trang_thai') == 'OK':
        out += [(b['gia'], b['loai']) for b in m.get('ban', []) if b.get('gia') is not None]
    out += [(f['gia'][s], 'Đại lý') for s in src
            if s not in METASEARCH and f.get('gia', {}).get(s) is not None]
    return out


def verify(cp, path):
    k = Kiem()
    wf = load_workbook(path)                    # công thức + hyperlink
    wv = load_workbook(path, data_only=True)    # giá trị đã recalc
    src = cp['meta']['nguon']

    # ---- 1. công thức đã tính ----
    thieu = []
    for ws in wf.worksheets:
        vs = wv[ws.title]
        for row in ws.iter_rows():
            for c in row:
                if not (isinstance(c.value, str) and c.value.startswith('=')):
                    continue
                v = vs[c.coordinate].value
                if v is None and ',"")' not in c.value:      # IFERROR(...,"") được phép rỗng
                    thieu.append('%s!%s' % (ws.title, c.coordinate))
                k(not (isinstance(v, str) and v in LOI_EXCEL),
                  'Ô %s!%s lỗi công thức: %s' % (ws.title, c.coordinate, v))
    if not k(not thieu, '%d ô công thức chưa có giá trị (chưa recalc?): %s'
             % (len(thieu), ', '.join(thieu[:5]))):
        return k                                 # chưa recalc thì kiểm tiếp vô nghĩa

    # ---- 2. Data ----
    ws, wsf = wv['Data'], wf['Data']
    hdrs = dong_tieu_de(ws)
    chang = [(P, leg) for P in cp['caps'] for leg in ('di', 've')]
    k(len(hdrs) == len(chang), 'Data: %d khối chuyến, checkpoint có %d' % (len(hdrs), len(chang)))
    log = cp.get('log', [])
    hang_data = {}                               # (cap, leg) -> [số dòng]
    for h, (P, leg) in zip(hdrs, chang):
        col = bang(ws, h)
        ngay = P['ngay_' + leg]
        rows = []
        r = h + 1
        while ws.cell(r, 1).value == P['ten']:
            rows.append(r); r += 1
        hang_data[(P['ten'], leg)] = rows
        k(len(rows) == len(P[leg]['chuyen']), 'Data %s %s: %d dòng, checkpoint %d chuyến'
          % (P['ten'], leg, len(rows), len(P[leg]['chuyen'])))
        for r, f in zip(rows, P[leg]['chuyen']):
            nhan = 'Data!%d (%s %s %s)' % (r, ngay, f['dep'], f['hang'])
            k(ws.cell(r, col['Hãng']).value == f['hang'] and hhmm(ws.cell(r, col['Giờ khởi hành']).value) == f['dep'],
              '%s: hãng/giờ không khớp checkpoint' % nhan)
            for s in src:
                v = ws.cell(r, col['Giá %s (1 người lớn)' % s]).value
                g = f.get('gia', {}).get(s)
                if g is None:
                    k(v == NA, '%s: %s phải là "%s", đang là %r' % (nhan, s, NA, v))
                    k(any(e.get('cap') == ngay and e.get('nguon') == s and f['dep'] in e.get('ghi_chu', '')
                          for e in log), '%s: ô %s "%s" không có dòng Log giải thích' % (nhan, s, NA))
                else:
                    k(gan(v, g), '%s: giá %s = %r, checkpoint = %s' % (nhan, s, v, g))
            c = chot(f, src)
            v = ws.cell(r, col['Giá chốt (nguồn ưu tiên)']).value
            k(gan(v, c) if c is not None else v == NA, '%s: Giá chốt = %r, đúng ra %r' % (nhan, v, c))
            g = [x for x in f.get('gia', {}).values() if x is not None]
            if 'Giá thấp nhất (mọi nguồn)' in col:
                v = ws.cell(r, col['Giá thấp nhất (mọi nguồn)']).value
                k(gan(v, min(g)) if g else v == NA, '%s: giá thấp nhất = %r, đúng ra %r' % (nhan, v, min(g) if g else NA))
                nm = ws.cell(r, col['Nguồn giá thấp nhất']).value or ''
                dung = {s for s in src if f.get('gia', {}).get(s) == min(g)} if g else {NA}
                k(set(nm.split(', ')) == dung, '%s: nguồn giá thấp nhất "%s", đúng ra "%s"'
                  % (nhan, nm, ', '.join(s for s in src if s in dung) or NA))
                v = ws.cell(r, col['Chênh lệch max %']).value
                if len(g) >= 2:
                    k(isinstance(v, (int, float)) and abs(v - (max(g) - min(g)) / min(g)) < 1e-6,
                      '%s: chênh lệch = %r' % (nhan, v))
            if f.get('link_chuyen'):
                hl = wsf.cell(r, col['Link chuyến này (trang đặt vé)']).hyperlink
                k(hl is not None and hl.target == f['link_chuyen'], '%s: link chuyến sai hoặc thiếu' % nhan)

    # Block C, tính lại bằng vét cạn
    reC = {}
    hc = dong_tieu_de(ws, 'Loại phương án')
    k(len(hc) == len(cp['caps']), 'Data: %d Block C, checkpoint %d cặp' % (len(hc), len(cp['caps'])))
    for h, P in zip(hc, cp['caps']):
        col = bang(ws, h)
        di = [(chot(f, src), f) for f in P['di']['chuyen'] if chot(f, src) is not None]
        ve = [(chot(f, src), f) for f in P['ve']['chuyen'] if chot(f, src) is not None]
        r = h + 1
        while ws.cell(r, 1).value in ('Rẻ nhất', 'Cùng hãng rẻ nhất'):
            ten = ws.cell(r, 1).value
            a, b, t = (ws.cell(r, col[x]).value for x in ('Giá đi', 'Giá về', 'Tổng theo Giá chốt'))
            k(gan(t, (a or 0) + (b or 0)), 'Block C %s %s: tổng %r != %r + %r' % (P['ten'], ten, t, a, b))
            if ten == 'Rẻ nhất':
                lo_di, lo_ve = min(x for x, _ in di), min(y for y, _ in ve)   # chỉ so giá: hoà giá không so dict
                k(gan(a, lo_di) and gan(b, lo_ve),
                  'Block C %s: Rẻ nhất %r/%r, vét cạn ra %r/%r' % (P['ten'], a, b, lo_di, lo_ve))
                reC[P['ten']] = (t, r)
            else:
                best = min(x + y for x, fx in di for y, fy in ve if fx['hang'] == fy['hang'])
                k(gan(t, best), 'Block C %s: Cùng hãng %r, vét cạn ra %r' % (P['ten'], t, best))
            r += 1
        k(P['ten'] in reC or not (di and ve), 'Block C %s: thiếu dòng Rẻ nhất' % P['ten'])

    # ---- 3. Summary ----
    ws, wsf = wv['Summary'], wf['Summary']
    h = dong_tieu_de(ws)[0]
    col = bang(ws, h)
    rows = [r for r in range(h + 1, ws.max_row + 1) if ws.cell(r, col['Tổng theo Giá chốt']).value is not None]
    k(len(rows) == len(reC), 'Summary: %d dòng, Block C có %d phương án Rẻ nhất' % (len(rows), len(reC)))
    tong = [ws.cell(r, col['Tổng theo Giá chốt']).value for r in rows]
    k(tong == sorted(tong), 'Summary: chưa sắp theo Tổng theo Giá chốt tăng dần')
    for r in rows:
        o = ws.cell(r, 1).value or ''
        # Ô ghi "<tên> (<ngày đi> → <ngày về>)"; tên cặp tự đặt có thể chứa " (" nên khớp cả chuỗi.
        P = next((p for p in cp['caps'] if o == '%s (%s → %s)' % (p['ten'], p['ngay_di'], p['ngay_ve'])), None)
        if not k(P is not None, 'Summary dòng %d: "%s" không khớp cặp nào trong checkpoint' % (r, o)):
            continue
        ten = P['ten']
        k(ten in reC and gan(ws.cell(r, col['Tổng theo Giá chốt']).value, reC[ten][0]),
          'Summary %s: tổng không khớp Block C' % ten)
        for h_link in ('Link chuyến đi', 'Link chuyến về'):
            hl = wsf.cell(r, col[h_link]).hyperlink
            k(hl is not None and hl.target, 'Summary %s: thiếu %s' % (ten, h_link))
        chon = []
        for leg, h_cf in (('di', 'Chuyến đi'), ('ve', 'Chuyến về')):
            txt = ws.cell(r, col[h_cf]).value or ''
            f = next((f for f in P[leg]['chuyen'] if txt == '%s %s, %s' % (f['hang'], f.get('sh') or '-', f['dep'])), None)
            k(f is not None, 'Summary %s: không tìm thấy chuyến "%s" trong checkpoint' % (ten, txt))
            chon.append((f, P[leg]))
        if all(f for f, _ in chon):
            bans = [ban_cua(f, c, src) for f, c in chon]
            v = ws.cell(r, col['Tổng mua được rẻ nhất']).value
            dung = sum(min(b)[0] for b in bans) if all(bans) else 'Chưa xác nhận'
            k(gan(v, dung) if all(bans) else v == dung, 'Summary %s: tổng mua được rẻ nhất %r, đúng ra %r' % (ten, v, dung))
            hang = [[g for g, l in b if l == 'Hãng hàng không'] for b in bans]
            v = ws.cell(r, col['Tổng nếu mua thẳng từ hãng']).value
            dung = sum(min(x) for x in hang) if all(hang) else 'Chưa xác nhận'
            k(gan(v, dung) if all(hang) else v == dung, 'Summary %s: tổng mua thẳng từ hãng %r, đúng ra %r' % (ten, v, dung))

    # ---- 4. Mua vé ----
    ws, wsf = wv['Mua vé'], wf['Mua vé']
    hs = dong_tieu_de(ws)
    if cp.get('de_xuat') and k(hs, 'Mua vé: không có dòng tiêu đề'):
        col = bang(ws, hs[0])
        thay = set()
        for r in range(hs[0] + 1, ws.max_row + 1):
            if not ws.cell(r, 1).value:
                continue
            ngay, sh, gia = (ws.cell(r, col[x]).value for x in ('Ngày bay', 'Chuyến', 'Giá (1 người lớn)'))
            nguon, dep = ws.cell(r, col['Nguồn dữ liệu']).value or '', hhmm(ws.cell(r, col['Giờ đi']).value)
            thay.add((ngay, dep, sh))
            if not isinstance(gia, (int, float)):
                continue
            hl = wsf.cell(r, col['Link mua']).hyperlink
            k(hl is not None and hl.target, 'Mua vé dòng %d: thiếu link mua' % r)
            k(bool(ws.cell(r, col['Thời điểm']).value), 'Mua vé dòng %d: thiếu thời điểm' % r)
            fs = [f for P in cp['caps'] for leg in ('di', 've') if P['ngay_' + leg] == ngay
                  for f in P[leg]['chuyen'] if f['dep'] == dep and (f.get('sh') or '-') == sh]
            if not k(bool(fs), 'Mua vé dòng %d: không thấy chuyến %s %s %s' % (r, ngay, dep, sh)):
                continue
            f = fs[0]
            if nguon == 'Trang đặt vé Google Flights':
                ok = any(b.get('gia') == gia for b in (f.get('mua') or {}).get('ban', []))
            else:
                ok = f.get('gia', {}).get(nguon.replace('Quét trực tiếp ', '')) == gia
            k(ok, 'Mua vé dòng %d: giá %r không khớp %s' % (r, gia, nguon))
        for e in cp['de_xuat']:
            ngay = '%s/%s/%s' % (e['ngay'][8:10], e['ngay'][5:7], e['ngay'][0:4])
            k((ngay, e['dep'], e['sh'] or '-') in thay,
              'Mua vé: thiếu chuyến đề xuất %s %s %s' % (ngay, e['dep'], e['sh']))
    return k


def main(cp_path, xlsx):
    with open(cp_path, encoding='utf-8') as fh:
        cp = json.load(fh)
    k = verify(cp, xlsx)
    print(json.dumps({'file': xlsx, 'so_kiem': k.n, 'so_loi': len(k.loi)}, ensure_ascii=False))
    for m in k.loi[:40]:
        print('LỖI: ' + m)
    if len(k.loi) > 40:
        print('... và %d lỗi khác' % (len(k.loi) - 40))
    return 1 if k.loi else 0


if __name__ == '__main__':
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    sys.exit(main(sys.argv[1], sys.argv[2]))
