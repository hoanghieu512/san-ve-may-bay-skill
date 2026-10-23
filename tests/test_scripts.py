# -*- coding: utf-8 -*-
"""Tests for the skill's Python scripts. Not packaged into the .skill file.

    python3 -m unittest discover -s tests -v

Workbook checks that need computed formula values run only when LibreOffice
(`soffice`) is on PATH; they are skipped otherwise.
"""
import datetime
import json
import os
import shutil
import subprocess
import sys
import tempfile
import unittest

sys.dont_write_bytecode = True   # keep __pycache__ out of the skill folder that gets zipped

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SCRIPTS = os.path.join(ROOT, 'skill', 've-may-bay-noi-dia', 'scripts')
sys.path.insert(0, SCRIPTS)

import merge_sources  # noqa: E402
import shortlist      # noqa: E402

EM_DASH = chr(0x2014)

# Captured live from Google Flights on 23/10/2026: the booking page of VJ 632 SGN-DAD 14/11/2026.
LIVE_TFS = ('CBwQAho_EgoyMDI2LTExLTE0Ih8KA1NHThIKMjAyNi0xMS0xNBoDREFEKgJWSjIDNjMyagcIARIDU0dO'
            'cgcIARIDREFEQAFIAXABggELCP___________wGYAQI')


def run(script, *args, check=True):
    p = subprocess.run([sys.executable, os.path.join(SCRIPTS, script)] + list(args),
                       capture_output=True, text=True,
                       env=dict(os.environ, PYTHONDONTWRITEBYTECODE='1'))
    if check and p.returncode != 0:
        raise AssertionError('%s failed (%d):\n%s\n%s' % (script, p.returncode, p.stdout, p.stderr))
    return p


def write(d, name, obj):
    with open(os.path.join(d, name), 'w', encoding='utf-8') as fh:
        json.dump(obj, fh, ensure_ascii=False)


def make_job(d, with_xac_nhan=True):
    """Two date pairs, GF (no flight numbers on one row) + Traveloka, per-leg capture times."""
    write(d, 'gf_di.json', {
        '_thu_thap': {'2026-11-14': '23/10/2026 10:01 (giờ VN)', '2026-11-15': '23/10/2026 10:02 (giờ VN)'},
        '2026-11-14': [['20:10', '21:30', 'Vietjet', 'VJ 632', 742061],
                       ['05:45', '07:15', 'Vietnam Airlines', 'VN 6030', 1127181],
                       ['05:50', '07:15', 'Sun PhuQuoc Airways', '9G 954', 906181],
                       ['09:00', '10:20', 'Bamboo Airways', '', 1500000]],
        '2026-11-15': [['20:10', '21:30', 'Vietjet', 'VJ 632', 800000]]})
    write(d, 'gf_ve.json', {
        '_thu_thap': {'2026-11-17': '23/10/2026 10:03 (giờ VN)'},
        '2026-11-17': [['18:00', '19:20', 'Vietjet', 'VJ 631', 1100000],
                       ['19:00', '20:20', 'Vietnam Airlines', 'VN 135', 1300000]]})
    write(d, 'tvl.json', {
        'di': {'_thu_thap': {'2026-11-14': '23/10/2026 10:40 (giờ VN)', '2026-11-15': '23/10/2026 10:41 (giờ VN)'},
               '2026-11-14': [['VJ-632', '20:10', '21:30', 'VietJet Air', 700000],
                              ['VN-6030', '05:45', '07:15', 'Vietnam Airlines', 1127181],
                              ['9G-954', '05:50', '07:15', 'SUN PhuQuoc AIRWAYS', 1500000]],
               '2026-11-15': [['VJ-632', '20:10', '21:30', 'VietJet Air', 790000]]},
        've': {'_thu_thap': {'2026-11-17': '23/10/2026 10:42 (giờ VN)'},
               '2026-11-17': [['VJ-631', '18:00', '19:20', 'VietJet Air', 1050000],
                              ['VN-135', '19:00', '20:20', 'Vietnam Airlines', 1290000]]}})
    job = {
        'meta': {'chang': 'SGN ⇄ DAD', 'khach': '1 người lớn', 'hang_ve': 'Phổ thông',
                 'loai': 'Chỉ bay thẳng'},
        'nguon': ['Google Flights', 'Traveloka'],
        'san_bay': {'di': ['SGN', 'DAD'], 've': ['DAD', 'SGN']},
        'raw': {'Google Flights': {'di': 'gf_di.json', 've': 'gf_ve.json',
                                   'cot': ['dep', 'arr', 'hang', 'sh', 'gia']},
                'Traveloka': {'file': 'tvl.json', 'cot': ['sh', 'dep', 'arr', 'hang', 'gia']}},
        'caps': [{'id': 1, 'ngay_di': '14/11/2026', 'ngay_ve': '17/11/2026', 'thu': 'T7 → T3'},
                 {'id': 2, 'ngay_di': '15/11/2026', 'ngay_ve': '17/11/2026', 'thu': 'CN → T3'}],
    }
    if with_xac_nhan:
        url = merge_sources.gf_booking_url('2026-11-14', 'SGN', 'DAD', 'VJ 632')
        write(d, 'gf_booking_raw.json', {'ket_qua': [
            {'url': url, 'sh': 'VJ 632', 'dep': '20:10', 'ngay': '2026-11-14', 'tu': 'SGN', 'den': 'DAD',
             'luc': '23/10/2026 11:22 (giờ VN)', 'trang_thai': 'OK', 'ly_do': '',
             'chuyen': 'Thứ 7, 14 thg 11 | 20:10 | – | 21:30 | Vietjet',
             'hanh_ly': ['Miễn phí 1 kiện hành lý xách tay', 'Có tính phí kiện hành lý ký gửi đầu tiên'],
             'ban': [{'nguoi_ban': 'Vietjet', 'loai': 'Hãng hàng không', 'gia': 742061,
                      'ghi_chu': 'Bao gồm phí tiện lợi 100.000 ₫'},
                     {'nguoi_ban': 'Jettzy', 'loai': 'Đại lý', 'gia': 679301, 'ghi_chu': ''}]},
            {'url': merge_sources.gf_booking_url('2026-11-17', 'DAD', 'SGN', 'VJ 631'), 'sh': 'VJ 631',
             'dep': '18:00', 'ngay': '2026-11-17', 'tu': 'DAD', 'den': 'SGN',
             'luc': '23/10/2026 11:23 (giờ VN)', 'trang_thai': 'OK', 'ly_do': '', 'chuyen': '',
             'hanh_ly': [], 'ban': [{'nguoi_ban': 'Vietjet', 'loai': 'Hãng hàng không',
                                     'gia': 1150000, 'ghi_chu': ''}]},
            {'url': merge_sources.gf_booking_url('2026-11-14', 'SGN', 'DAD', '9G 954'), 'sh': '9G 954',
             'dep': '05:50', 'ngay': '2026-11-14', 'tu': 'SGN', 'den': 'DAD',
             'luc': '23/10/2026 11:23 (giờ VN)', 'trang_thai': 'Không khả dụng',
             'ly_do': 'Google Flights báo hành trình không còn', 'chuyen': '', 'hanh_ly': [], 'ban': []},
        ], 'con_lai': [], 'dung': ''})
        job['xac_nhan'] = ['gf_booking_raw.json']
    write(d, 'job.json', job)
    return os.path.join(d, 'job.json')


def header_row(ws):
    return next(r for r in range(1, 30) if ws.cell(r, 1).value == 'Cặp ngày')


class TestShortlist(unittest.TestCase):
    NG = ['A', 'B']

    def f(self, hang, a=None, b=None, dep='06:00'):
        g = {}
        if a is not None:
            g['A'] = a
        if b is not None:
            g['B'] = b
        return {'hang': hang, 'dep': dep, 'gia': g}

    def test_gia_chot_follows_priority_not_minimum(self):
        self.assertEqual(shortlist.gia_chot(self.f('X', 200, 100), self.NG), 200)
        self.assertEqual(shortlist.gia_chot(self.f('X', None, 100), self.NG), 100)
        self.assertIsNone(shortlist.gia_chot(self.f('X'), self.NG))

    def test_phuong_an_cheapest_and_same_airline(self):
        cap = {'di': {'chuyen': [self.f('X', 100), self.f('Y', 90)]},
               've': {'chuyen': [self.f('X', 100), self.f('Z', 50)]}}
        pa = {t: (a[0], b[0]) for t, a, b, _ in shortlist.phuong_an(cap, self.NG)}
        self.assertEqual(pa['Rẻ nhất'], (1, 1))
        self.assertEqual(pa['Cùng hãng rẻ nhất'], (0, 0))

    def test_de_xuat_includes_block_c_and_top_n(self):
        di = [self.f('X', p, dep='%02d:00' % i) for i, p in enumerate([500, 100, 300, 200, 400])]
        cap = {'di': {'chuyen': di}, 've': {'chuyen': [self.f('X', 50)]}}
        dx = shortlist.de_xuat(cap, self.NG, top=2)
        self.assertEqual(sorted(dx['di']), [1, 3])          # 100 (Block C + top) and 200 (top)
        self.assertIn('Rẻ nhất', dx['di'][1])
        self.assertEqual(sorted(dx['ve']), [0])


class TestMerge(unittest.TestCase):
    def setUp(self):
        self.d = tempfile.mkdtemp()

    def tearDown(self):
        shutil.rmtree(self.d)

    def merge(self, *extra, **kw):
        job = make_job(self.d, **kw)
        cp_path = os.path.join(self.d, 'cp.json')
        p = run('merge_sources.py', job, cp_path, *extra)
        with open(cp_path, encoding='utf-8') as fh:
            return json.load(fh), p

    def test_booking_url_matches_live_google_flights(self):
        url = merge_sources.gf_booking_url('2026-11-14', 'SGN', 'DAD', 'VJ 632')
        self.assertIn('tfs=' + LIVE_TFS + '&', url)
        self.assertEqual(url, merge_sources.gf_booking_url('2026-11-14', 'SGN', 'DAD', 'VJ-632'))
        self.assertIsNone(merge_sources.gf_booking_url('2026-11-14', 'SGN', 'DAD', ''))

    def test_run_timestamp_is_vietnam_time(self):
        cp, _ = self.merge()
        ts = datetime.datetime.strptime(cp['meta']['thu_thap'][:16], '%d/%m/%Y %H:%M')
        vn_now = datetime.datetime.now(datetime.timezone.utc).replace(tzinfo=None) + datetime.timedelta(hours=7)
        self.assertLess(abs((vn_now - ts).total_seconds()), 180)

    def test_search_link_is_well_formed(self):
        cp, _ = self.merge()
        gf = cp['caps'][0]['di']['link']['Google Flights']
        self.assertIn('from%20SGN%20to%20DAD%20on%202026-11-14', gf)

    def test_per_leg_capture_time_and_underscore_keys(self):
        cp, _ = self.merge()
        c1 = cp['caps'][0]
        self.assertEqual(c1['di']['thu_thap'], {'Google Flights': '23/10/2026 10:01 (giờ VN)',
                                                'Traveloka': '23/10/2026 10:40 (giờ VN)'})
        self.assertEqual(c1['ve']['thu_thap']['Traveloka'], '23/10/2026 10:42 (giờ VN)')
        self.assertFalse(any(f['dep'].startswith('_') for f in c1['di']['chuyen']))

    def test_flight_link_and_number(self):
        cp, _ = self.merge()
        di = {f['sh']: f for f in cp['caps'][0]['di']['chuyen']}
        self.assertIn('tfs=' + LIVE_TFS, di['VJ 632']['link_chuyen'])
        bamboo = [f for f in cp['caps'][0]['di']['chuyen'] if f['hang'] == 'Bamboo Airways'][0]
        self.assertEqual(bamboo['link_chuyen'], '')

    def test_de_xuat_list_and_cli_output(self):
        cp, p = self.merge('--de-xuat')
        keys = {(x['ngay'], x['sh']) for x in cp['de_xuat']}
        self.assertIn(('2026-11-14', 'VJ 632'), keys)
        self.assertIn(('2026-11-17', 'VJ 631'), keys)
        # 17/11 return leg is shared by both pairs: listed once
        self.assertEqual(sum(1 for x in cp['de_xuat'] if x['ngay'] == '2026-11-17' and x['sh'] == 'VJ 631'), 1)
        start = p.stdout.index('[')
        flights = json.loads(p.stdout[start:p.stdout.rindex(']') + 1])
        self.assertTrue(all(x['url'] for x in flights))
        self.assertTrue(all(set(x) == {'url', 'sh', 'dep', 'ngay', 'tu', 'den'} for x in flights))

    def test_xac_nhan_attached_to_every_matching_flight(self):
        cp, _ = self.merge()
        vj = [f for f in cp['caps'][0]['di']['chuyen'] if f['sh'] == 'VJ 632'][0]
        self.assertEqual(vj['mua']['trang_thai'], 'OK')
        self.assertEqual([b['gia'] for b in vj['mua']['ban']], [742061, 679301])
        ve1 = [f for f in cp['caps'][0]['ve']['chuyen'] if f['sh'] == 'VJ 631'][0]
        ve2 = [f for f in cp['caps'][1]['ve']['chuyen'] if f['sh'] == 'VJ 631'][0]
        self.assertEqual(ve1['mua'], ve2['mua'])

    def test_xac_nhan_logs(self):
        cp, p = self.merge()
        txt = json.dumps(cp['log'], ensure_ascii=False)
        self.assertIn('9G 954', txt)                       # failed confirmation logged
        self.assertIn('Chưa xác nhận nơi mua', txt)       # shortlisted but no result
        changed = [e for e in cp['log'] if 'đổi sau khi quét' in e['ly_do']]
        # return leg VJ 631 changed 1,100,000 -> 1,150,000 since the scan; VJ 632 did not change
        self.assertEqual([e['cap'] for e in changed], ['17/11/2026'])
        self.assertIn('VJ 631', changed[0]['ghi_chu'])

    def test_missing_xac_nhan_warns(self):
        cp, p = self.merge(with_xac_nhan=False)
        self.assertIn('xác nhận nơi mua', p.stderr)
        self.assertIn('Chưa chạy bước xác nhận nơi mua', json.dumps(cp['log'], ensure_ascii=False))

    def test_flight_number_nbsp_normalized(self):
        # Google Flights renders "VJ\u00a0636"; the checkpoint must hold a plain space.
        job = make_job(self.d)
        with open(os.path.join(self.d, 'gf_di.json'), encoding='utf-8') as fh:
            gf = json.load(fh)
        gf['2026-11-14'][0][3] = 'VJ\u00a0632'
        write(self.d, 'gf_di.json', gf)
        cp_path = os.path.join(self.d, 'cp.json')
        run('merge_sources.py', job, cp_path)
        with open(cp_path, encoding='utf-8') as fh:
            cp = json.load(fh)
        shs = [f['sh'] for c in cp['caps'] for leg in ('di', 've') for f in c[leg]['chuyen']]
        self.assertIn('VJ 632', shs)
        self.assertFalse(any('\u00a0' in x for x in shs))

    def test_xac_nhan_prefers_latest_ok_over_later_failure(self):
        job = make_job(self.d)
        with open(os.path.join(self.d, 'gf_booking_raw.json'), encoding='utf-8') as fh:
            first = json.load(fh)
        first['dung'] = 'Google chặn (captcha/consent) tại https://www.google.com/sorry/index'
        write(self.d, 'gf_booking_raw.json', first)
        retry = dict(first['ket_qua'][0], trang_thai='Lỗi', ly_do='Trang đặt vé không nạp xong', ban=[])
        write(self.d, 'gf_booking_raw_2.json', {'ket_qua': [retry], 'con_lai': [], 'dung': ''})
        with open(job, encoding='utf-8') as fh:
            j = json.load(fh)
        j['xac_nhan'] = ['gf_booking_raw.json', 'gf_booking_raw_2.json']
        write(self.d, 'job.json', j)
        cp_path = os.path.join(self.d, 'cp.json')
        run('merge_sources.py', job, cp_path)
        with open(cp_path, encoding='utf-8') as fh:
            cp = json.load(fh)
        vj = [f for f in cp['caps'][0]['di']['chuyen'] if f['sh'] == 'VJ 632'][0]
        self.assertEqual(vj['mua']['trang_thai'], 'OK')            # earlier OK beats later failure
        blocked = [e for e in cp['log'] if 'Bị chặn' in e['ly_do']]
        self.assertTrue(blocked)
        self.assertTrue(all(e['trang_thai'] != 'Lỗi' for e in blocked))   # superseded by a later file

    def test_no_em_dash_in_generated_text(self):
        cp, _ = self.merge()
        self.assertNotIn(EM_DASH, json.dumps(cp, ensure_ascii=False))


class TestWorkbook(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.d = tempfile.mkdtemp()
        job = make_job(cls.d)
        cls.cp = os.path.join(cls.d, 'cp.json')
        cls.xlsx = os.path.join(cls.d, 'out.xlsx')
        run('merge_sources.py', job, cls.cp)
        run('build_workbook.py', cls.cp, cls.xlsx)
        from openpyxl import load_workbook
        cls.wb = load_workbook(cls.xlsx)

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.d)

    def test_sheet_order(self):
        self.assertEqual(self.wb.sheetnames, ['Summary', 'Mua vé', 'Data', 'Log'])

    def test_summary_links_both_legs(self):
        ws = self.wb['Summary']
        h = header_row(ws)
        hdr = [c.value for c in ws[h]]
        for name in ('Link chuyến đi', 'Link chuyến về'):
            col = hdr.index(name) + 1
            for r in range(h + 1, h + 3):
                self.assertTrue(ws.cell(r, col).hyperlink, '%s row %d has no hyperlink' % (name, r))

    def test_summary_confirmed_totals(self):
        ws = self.wb['Summary']
        h = header_row(ws)
        hdr = [c.value for c in ws[h]]
        rows = {ws.cell(r, 1).value: [c.value for c in ws[r]] for r in range(h + 1, h + 3)}
        c1 = [v for k, v in rows.items() if k.startswith('Cặp 1')][0]
        # cheapest seller per leg: Jettzy 679,301 (GF booking page) + Traveloka 1,050,000 (direct scan)
        self.assertEqual(c1[hdr.index('Tổng mua được rẻ nhất')], 679301 + 1050000)
        self.assertEqual(c1[hdr.index('Nơi mua rẻ nhất (đi / về)')], 'Jettzy (đại lý) / Traveloka (đại lý)')
        self.assertEqual(c1[hdr.index('Tổng nếu mua thẳng từ hãng')], 742061 + 1150000)

    def test_mua_ve_rows(self):
        ws = self.wb['Mua vé']
        hdr_row = header_row(ws)
        hdr = [c.value for c in ws[hdr_row]]
        rows = [[c.value for c in ws[r]] for r in range(hdr_row + 1, ws.max_row + 1)
                if ws.cell(r, 1).value]
        ban = [r for r in rows if r[hdr.index('Chuyến')] == 'VJ 632' and r[hdr.index('Ngày bay')] == '14/11/2026']
        sellers = [r[hdr.index('Người bán')] for r in ban]
        self.assertEqual(sellers[:2], ['Jettzy', 'Traveloka'])   # sorted by price, 679,301 < 700,000
        self.assertIn('Vietjet', sellers)
        for r in ban:
            self.assertTrue(r[hdr.index('Link mua')])
            self.assertTrue(r[hdr.index('Thời điểm')])
        failed = [r for r in rows if r[hdr.index('Chuyến')] == '9G 954']
        self.assertTrue(failed and failed[0][hdr.index('Người bán')] == 'Không khả dụng')

    def test_data_has_min_source_and_flight_link(self):
        ws = self.wb['Data']
        hdr_row = header_row(ws)
        hdr = [c.value for c in ws[hdr_row]]
        self.assertIn('Nguồn giá thấp nhất', hdr)
        self.assertIn('Link chuyến này (trang đặt vé)', hdr)
        self.assertTrue(any('1 người lớn' in (h or '') for h in hdr))

    def test_no_em_dash_anywhere(self):
        for ws in self.wb.worksheets:
            for row in ws.iter_rows():
                for c in row:
                    if isinstance(c.value, str):
                        self.assertNotIn(EM_DASH, c.value, '%s!%s' % (ws.title, c.coordinate))

    @unittest.skipUnless(shutil.which('soffice'), 'LibreOffice not installed')
    def test_verify_passes_then_catches_tampering(self):
        out = os.path.join(self.d, 'rc')
        subprocess.run(['soffice', '--headless', '--calc', '--convert-to', 'xlsx', '--outdir', out,
                        self.xlsx], capture_output=True, timeout=180, check=True)
        good = os.path.join(out, 'out.xlsx')
        p = run('verify_workbook.py', self.cp, good, check=False)
        self.assertEqual(p.returncode, 0, p.stdout + p.stderr)

        not_recalc = run('verify_workbook.py', self.cp, self.xlsx, check=False)
        self.assertNotEqual(not_recalc.returncode, 0)
        self.assertIn('recalc', not_recalc.stdout + not_recalc.stderr)

        from openpyxl import load_workbook
        wb = load_workbook(good)
        ws = wb['Data']
        hdr_row = header_row(ws)
        ws.cell(hdr_row + 1, 7).value = 1          # tamper a source price
        bad = os.path.join(self.d, 'bad.xlsx')
        wb.save(bad)
        subprocess.run(['soffice', '--headless', '--calc', '--convert-to', 'xlsx', '--outdir',
                        os.path.join(self.d, 'rc2'), bad], capture_output=True, timeout=180, check=True)
        p = run('verify_workbook.py', self.cp, os.path.join(self.d, 'rc2', 'bad.xlsx'), check=False)
        self.assertNotEqual(p.returncode, 0)


def recalc(xlsx, outdir):
    subprocess.run(['soffice', '--headless', '--calc', '--convert-to', 'xlsx', '--outdir', outdir, xlsx],
                   capture_output=True, timeout=180, check=True)
    return os.path.join(outdir, os.path.basename(xlsx))


class TestBuildEdgeCases(unittest.TestCase):
    def test_empty_de_xuat_is_not_reported_as_old_checkpoint(self):
        d = tempfile.mkdtemp()
        try:
            cp = {'meta': {'chang': 'X', 'khach': '1 người lớn', 'hang_ve': 'Phổ thông', 'loai': 'Bay thẳng',
                           'thu_thap': '23/10/2026 10:00 (giờ VN)', 'nguon': ['Google Flights']},
                  'caps': [{'id': 1, 'ten': 'Cặp 1', 'ngay_di': '14/11/2026', 'ngay_ve': '17/11/2026', 'thu': '',
                            'di': {'link': {}, 'chuyen': []}, 've': {'link': {}, 'chuyen': []}}],
                  'de_xuat': [], 'log': []}
            write(d, 'cp.json', cp)
            run('build_workbook.py', os.path.join(d, 'cp.json'), os.path.join(d, 'out.xlsx'))
            from openpyxl import load_workbook
            ws = load_workbook(os.path.join(d, 'out.xlsx'))['Mua vé']
            txt = ' '.join(str(c.value) for row in ws.iter_rows() for c in row if c.value)
            self.assertNotIn('bản 1.5.0', txt)
            self.assertIn('Không có chuyến nào', txt)
        finally:
            shutil.rmtree(d)


@unittest.skipUnless(shutil.which('soffice'), 'LibreOffice not installed')
class TestVerifyEdgeCases(unittest.TestCase):
    def setUp(self):
        self.d = tempfile.mkdtemp()

    def tearDown(self):
        shutil.rmtree(self.d)

    def build(self, edit_job=None, edit_gf=None):
        job = make_job(self.d, with_xac_nhan=False)
        if edit_gf:
            with open(os.path.join(self.d, 'gf_di.json'), encoding='utf-8') as fh:
                gf = json.load(fh)
            edit_gf(gf)
            write(self.d, 'gf_di.json', gf)
        if edit_job:
            with open(job, encoding='utf-8') as fh:
                j = json.load(fh)
            edit_job(j)
            write(self.d, 'job.json', j)
        cp, xlsx = os.path.join(self.d, 'cp.json'), os.path.join(self.d, 'out.xlsx')
        run('merge_sources.py', job, cp)
        run('build_workbook.py', cp, xlsx)
        return cp, recalc(xlsx, os.path.join(self.d, 'rc'))

    def test_cap_name_with_parenthesis(self):
        def ten(j):
            j['caps'][0]['ten'] = 'Cặp 1 (Tết)'
        cp, good = self.build(edit_job=ten)
        p = run('verify_workbook.py', cp, good, check=False)
        self.assertEqual(p.returncode, 0, p.stdout + p.stderr)

    def test_min_source_must_list_every_tied_source(self):
        def tie(gf):
            gf['2026-11-14'][0][4] = 700000                 # VJ 632: Google Flights == Traveloka
        cp, good = self.build(edit_gf=tie)
        self.assertEqual(run('verify_workbook.py', cp, good, check=False).returncode, 0)
        from openpyxl import load_workbook
        wb = load_workbook(good)
        ws = wb['Data']
        h = header_row(ws)
        col = {ws.cell(h, c).value: c for c in range(1, ws.max_column + 1)}
        r = next(r for r in range(h + 1, ws.max_row + 1) if ws.cell(r, col['Số hiệu chuyến']).value == 'VJ 632')
        self.assertEqual(ws.cell(r, col['Nguồn giá thấp nhất']).value, 'Google Flights, Traveloka')
        ws.cell(r, col['Nguồn giá thấp nhất']).value = 'Traveloka'      # drop one tied source
        bad = os.path.join(self.d, 'bad.xlsx')
        wb.save(bad)
        p = run('verify_workbook.py', cp, recalc(bad, os.path.join(self.d, 'rc2')), check=False)
        self.assertNotEqual(p.returncode, 0)
        self.assertIn('nguồn giá thấp nhất', p.stdout)


def node_env():
    """Env where require('playwright') resolves, or None if node/playwright is missing."""
    if not shutil.which('node'):
        return None
    env = dict(os.environ)
    root = subprocess.run(['npm', 'root', '-g'], capture_output=True, text=True).stdout.strip()         if shutil.which('npm') else ''
    if root:
        env['NODE_PATH'] = os.pathsep.join(x for x in (env.get('NODE_PATH'), root) if x)
    ok = subprocess.run(['node', '-e', "require('playwright')"], capture_output=True, env=env)
    return env if ok.returncode == 0 else None


class TestGfBookingOffline(unittest.TestCase):
    @unittest.skipUnless(node_env(), 'node + playwright not available')
    def test_offline_scenarios(self):
        p = subprocess.run(['node', os.path.join(ROOT, 'tests', 'gf_booking_offline.js'),
                            os.path.join(SCRIPTS, 'gf_booking.js')],
                           capture_output=True, text=True, env=node_env(), timeout=180)
        self.assertEqual(p.returncode, 0, p.stdout + p.stderr)
        self.assertEqual(p.stdout.count('PASS'), 6, p.stdout)


class TestVerifyTies(unittest.TestCase):
    @unittest.skipUnless(shutil.which('soffice'), 'LibreOffice not installed')
    def test_two_flights_same_price(self):
        d = tempfile.mkdtemp()
        try:
            job = make_job(d, with_xac_nhan=False)
            with open(os.path.join(d, 'gf_di.json'), encoding='utf-8') as fh:
                gf = json.load(fh)
            gf['2026-11-14'][1][4] = gf['2026-11-14'][0][4]       # VN 6030 priced like VJ 632
            write(d, 'gf_di.json', gf)
            cp, xlsx = os.path.join(d, 'cp.json'), os.path.join(d, 'out.xlsx')
            run('merge_sources.py', job, cp)
            run('build_workbook.py', cp, xlsx)
            subprocess.run(['soffice', '--headless', '--calc', '--convert-to', 'xlsx', '--outdir',
                            os.path.join(d, 'rc'), xlsx], capture_output=True, timeout=180, check=True)
            p = run('verify_workbook.py', cp, os.path.join(d, 'rc', 'out.xlsx'), check=False)
            self.assertEqual(p.returncode, 0, p.stdout + p.stderr)
        finally:
            shutil.rmtree(d)


if __name__ == '__main__':
    unittest.main()
